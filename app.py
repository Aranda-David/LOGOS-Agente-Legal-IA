import io
import html
import requests
import streamlit as st
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Importamos el extractor modular y el grafo
from src.document_loader import cargar_texto_desde_archivo
from src.graph import app_grafo

# Configuración de la página
st.set_page_config(
    page_title="LOGOS - Agente IA Legal",
    page_icon="⚖️",
    layout="wide"
)


# --- FUNCIONES DE UTILIDAD Y COMPROBACIÓN ---
def comprobar_estado_sistema():
    """Comprueba si el servidor local de Ollama está accesible."""
    try:
        response = requests.get("http://localhost:11434/api/tags", timeout=2)
        if response.status_code == 200:
            return True
    except Exception:
        pass
    return False


# --- FUNCIÓN PARA GENERAR DOCUMENTO WORD (.DOCX) ---
def crear_documento_word(texto_dictamen: str, titulo: str = "LOGOS — Dictamen / Documento Jurídico") -> io.BytesIO:
    """Genera un archivo .docx en memoria a partir del texto."""
    doc = Document()
    doc.add_heading(titulo, level=1)

    for line in texto_dictamen.split("\n"):
        if line.strip():
            doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# --- FUNCIÓN PARA GENERAR DOCUMENTO PDF (.PDF) ---
def crear_documento_pdf(texto_dictamen: str, titulo: str = "LOGOS — Dictamen / Documento Jurídico") -> io.BytesIO:
    """Genera un archivo PDF profesional en memoria a partir del texto."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        'TituloDictamen',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        spaceAfter=12
    )
    estilo_cuerpo = ParagraphStyle(
        'CuerpoDictamen',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        spaceAfter=8
    )

    story = []
    story.append(Paragraph(titulo, estilo_titulo))
    story.append(Spacer(1, 12))

    for linea in texto_dictamen.split("\n"):
        linea_limpia = linea.strip()
        if linea_limpia:
            linea_escapada = html.escape(linea_limpia)
            story.append(Paragraph(linea_escapada, estilo_cuerpo))
        else:
            story.append(Spacer(1, 6))

    doc.build(story)
    buffer.seek(0)
    return buffer


# Estilo y título principal
st.title("⚖️ LOGOS — Agente Legal Inteligente")
st.caption("Sistema local de auditoría y redacción jurídica impulsado por LangGraph y Ollama (RGPD Compliant)")

# --- BARRA LATERAL (SIDEBAR) ---
with st.sidebar:
    st.image("https://img.icons8.com/scales", width=80)
    st.header("⚙️ Panel de Control")
    st.write("Configura la sesión o adjunta documentación.")

    # 1. Indicador de estado de Ollama en tiempo real
    ollama_activo = comprobar_estado_sistema()
    if ollama_activo:
        st.markdown("🟢 **Motor Local (Ollama):** En línea")
    else:
        st.markdown("🔴 **Motor Local (Ollama):** Desconectado")

    st.divider()

    modo_analisis = st.selectbox(
        "Modo de trabajo preferido:",
        ["Enrutado Automático (IA)", "Auditoría de Riesgos", "Redacción Jurídica", "Consulta General"]
    )

    # 3. Selector manual de Jurisdicción opcional
    jurisdiccion_manual = st.selectbox(
        "Forzar Jurisdicción (Opcional):",
        ["Automática (IA)", "Penal", "Civil", "Laboral", "Mercantil", "Administrativo"]
    )

    st.divider()

    st.subheader("📄 Cargar Documento Legal")
    archivo_adjunto = st.file_uploader(
        "Adjunta tu documento (PDF, DOCX, RTF, TXT, MD)",
        type=["pdf", "docx", "rtf", "txt", "md"]
    )

    texto_documento_extraido = ""
    if archivo_adjunto is not None:
        texto_documento_extraido = cargar_texto_desde_archivo(archivo_adjunto)
        if texto_documento_extraido:
            st.success(f"✅ '{archivo_adjunto.name}' procesado correctamente.")
        else:
            st.error("⚠️ No se pudo extraer texto del archivo adjunto.")

    st.divider()

    # Inicializar historial si no existe para el bloque lateral de exportación
    if "mensajes" not in st.session_state:
        st.session_state.mensajes = []

    # 2. Exportación de sesión completa en la barra lateral si hay mensajes
    if st.session_state.mensajes:
        st.subheader("📦 Archivar Sesión")
        texto_historial_completo = "\n\n".join(
            [f"[{m['rol'].upper()}]: {m['contenido']}" for m in st.session_state.mensajes])

        col_exp1, col_exp2 = st.columns(2)
        with col_exp1:
            buf_w_sesion = crear_documento_word(texto_historial_completo, "LOGOS — Historial de Sesión")
            st.download_button("📥 Word", data=buf_w_sesion, file_name="Sesion_LOGOS.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               key="dl_sesion_w")
        with col_exp2:
            buf_p_sesion = crear_documento_pdf(texto_historial_completo, "LOGOS — Historial de Sesión")
            st.download_button("📄 PDF", data=buf_p_sesion, file_name="Sesion_LOGOS.pdf", mime="application/pdf",
                               key="dl_sesion_p")

        st.divider()

    if st.button("🗑️ Limpiar Conversación"):
        st.session_state.mensajes = []
        st.rerun()

st.divider()

# --- CHAT PRINCIPAL ---

# Mostrar historial con contexto RAG opcional y botones de descarga condicionales
for i, msg in enumerate(st.session_state.mensajes):
    with st.chat_message(msg["rol"]):
        st.markdown(msg["contenido"])

        if msg.get("contexto_rag") and not msg.get("es_conversacional", False):
            with st.expander("🔍 Ver Contexto RAG Recuperado"):
                st.markdown(msg["contexto_rag"])

        # Renderizar botones de descarga SOLO si NO es un mensaje conversacional
        if msg["rol"] == "assistant" and not msg.get("es_conversacional", False):
            col1, col2, _ = st.columns([1, 1, 3])
            with col1:
                buffer_word = crear_documento_word(msg["contenido"])
                st.download_button(
                    label="📥 Descargar Word (.docx)",
                    data=buffer_word,
                    file_name=f"Dictamen_LOGOS_{i}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"dl_word_{i}"
                )
            with col2:
                buffer_pdf = crear_documento_pdf(msg["contenido"])
                st.download_button(
                    label="📄 Descargar PDF (.pdf)",
                    data=buffer_pdf,
                    file_name=f"Dictamen_LOGOS_{i}.pdf",
                    mime="application/pdf",
                    key=f"dl_pdf_{i}"
                )

# Entrada del usuario
if peticion := st.chat_input("Escribe tu instrucción legal aquí..."):

    st.session_state.mensajes.append({"rol": "user", "contenido": peticion})
    with st.chat_message("user"):
        st.markdown(peticion)

    with st.chat_message("assistant"):
        # Estado estructurado inicial compatible con el grafo
        estado_inicial = {
            "peticion_usuario": peticion,
            "texto_documento": texto_documento_extraido if texto_documento_extraido else None,
            "tipo_tarea": None,
            "jurisdicciones": [],
            "contexto_rag": None,
            "borrador_respuesta": None,
            "respuesta_final": None
        }

        # Forzar el modo de análisis si no está en automático
        if modo_analisis == "Auditoría de Riesgos":
            estado_inicial["tipo_tarea"] = "auditoria"
        elif modo_analisis == "Redacción Jurídica":
            estado_inicial["tipo_tarea"] = "redaccion"
        elif modo_analisis == "Consulta General":
            estado_inicial["tipo_tarea"] = "general"

        # Forzar jurisdicción manual si se ha seleccionado en el panel
        if jurisdiccion_manual != "Automática (IA)":
            estado_inicial["jurisdicciones"] = [jurisdiccion_manual.lower()]

        # Contenedor mutable para capturar el contexto y el tipo de tarea detectada durante el streaming
        datos_ejecucion = {"contexto_rag": "", "tipo_tarea": ""}


        def generar_streaming():
            for output in app_grafo.stream(estado_inicial):
                for nodo_nombre, valor in output.items():
                    if isinstance(valor, dict):
                        if "tipo_tarea" in valor and valor["tipo_tarea"]:
                            datos_ejecucion["tipo_tarea"] = valor["tipo_tarea"]
                        if "contexto_rag" in valor and valor["contexto_rag"]:
                            datos_ejecucion["contexto_rag"] = valor["contexto_rag"]
                        if "respuesta_final" in valor and valor["respuesta_final"]:
                            yield valor["respuesta_final"]


        # Streaming en vivo del resultado final tras el nodo verificador
        respuesta_completa = st.write_stream(generar_streaming())

        contexto_rag_utilizado = datos_ejecucion["contexto_rag"]
        tipo_tarea_ejecutada = datos_ejecucion["tipo_tarea"]

        # Bandera para saber si fue charla coloquial/saludo
        es_conversacional = (tipo_tarea_ejecutada == "conversacional")

        # Desplegable para ver el contexto RAG consultado solo si aplica
        if contexto_rag_utilizado and not es_conversacional:
            with st.expander("🔍 Ver Contexto RAG Recuperado"):
                st.markdown(contexto_rag_utilizado)

        # Botones de descarga inmediatos SOLO si no es conversacional
        if not es_conversacional:
            col1, col2, _ = st.columns([1, 1, 3])
            with col1:
                buffer_word = crear_documento_word(respuesta_completa)
                st.download_button(
                    label="📥 Descargar Word (.docx)",
                    data=buffer_word,
                    file_name="Dictamen_LOGOS.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key="dl_word_nuevo"
                )
            with col2:
                buffer_pdf = crear_documento_pdf(respuesta_completa)
                st.download_button(
                    label="📄 Descargar PDF (.pdf)",
                    data=buffer_pdf,
                    file_name="Dictamen_LOGOS.pdf",
                    mime="application/pdf",
                    key="dl_pdf_nuevo"
                )

        st.session_state.mensajes.append({
            "rol": "assistant",
            "contenido": respuesta_completa,
            "contexto_rag": contexto_rag_utilizado,
            "es_conversacional": es_conversacional
        })